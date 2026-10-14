from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
EVIDENCE = BASE / "evidencias"
OUTPUT = BASE / "CampusFix_Actividad_Evaluable_06_Base_de_Datos_Cloud.docx"
DIAGRAM = BASE / "diagrama_er_campusfix.png"

NAVY = "17365D"
TEAL = "009C95"
PALE = "EAF2F8"
LIGHT = "F7F9FB"
BORDER = "D9D9D9"
GRAY = "5C6B7A"


def font(path_candidates, size):
    for path in path_candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def build_er_diagram():
    canvas = Image.new("RGB", (1800, 1080), "white")
    draw = ImageDraw.Draw(canvas)
    bold = font([r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\calibrib.ttf"], 36)
    body = font([r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\calibri.ttf"], 24)
    label = font([r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\calibrib.ttf"], 25)

    entities = {
        "USERS": (80, 90, 700, 430, ["uid  string  PK", "displayName  string", "email  string", "role  string", "active  boolean", "createdAt  timestamp"]),
        "CATEGORIES": (1100, 90, 1720, 390, ["categoryId  string  PK", "name  string", "description  string", "active  boolean", "order  number"]),
        "REPORTS": (520, 610, 1280, 1020, ["reportId  string  PK", "reporterId  string  FK", "categoryId  string  FK", "assigneeId  string or null", "title and description  string", "priority and status  string", "location  map", "evidencePath  string or null", "createdAt and updatedAt  timestamp"]),
        "HISTORY": (80, 650, 440, 1010, ["eventId  string  PK", "actorId  string  FK", "type  string", "previousStatus  string or null", "newStatus  string or null", "comment  string or null", "createdAt  timestamp"]),
    }

    def box(name, spec):
        x1, y1, x2, y2, fields = spec
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill="white", outline="#9FB3C8", width=4)
        draw.rounded_rectangle((x1, y1, x2, y1 + 68), radius=20, fill="#17365D")
        draw.rectangle((x1, y1 + 48, x2, y1 + 68), fill="#17365D")
        draw.text((x1 + 26, y1 + 15), name, fill="white", font=bold)
        y = y1 + 92
        for item in fields:
            draw.text((x1 + 28, y), item, fill="#26384A", font=body)
            y += 34

    for name, spec in entities.items():
        box(name, spec)

    # Relationship connectors are intentionally routed through open space.
    draw.line((390, 430, 650, 610), fill="#009C95", width=6)
    draw.polygon([(650, 610), (628, 575), (612, 600)], fill="#009C95")
    draw.text((370, 500), "crea  1 a N", fill="#006D68", font=label)

    draw.line((1410, 390, 1160, 610), fill="#009C95", width=6)
    draw.polygon([(1160, 610), (1188, 579), (1200, 607)], fill="#009C95")
    draw.text((1275, 475), "clasifica  1 a N", fill="#006D68", font=label)

    draw.line((440, 830, 520, 830), fill="#009C95", width=6)
    draw.polygon([(520, 830), (488, 812), (488, 848)], fill="#009C95")
    draw.text((445, 785), "N a 1", fill="#006D68", font=label)

    draw.text((650, 520), "Cada reporte pertenece a un usuario y una categoria", fill="#17365D", font=label)
    draw.text((650, 560), "y conserva un historial inmutable de eventos", fill="#17365D", font=label)
    canvas.save(DIAGRAM, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row_values):
            cells[idx].text = str(text)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            if row_index % 2 == 1:
                set_cell_shading(cells[idx], PALE)
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(32, 45, 58)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(92, 107, 122)


def add_image(doc, filename, width, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(EVIDENCE / filename), width=Inches(width))
    caption(doc, caption_text)


def add_two_images(doc, left_file, left_caption, right_file, right_caption):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (filename, label) in enumerate(((left_file, left_caption), (right_file, right_caption))):
        cell = table.cell(0, idx)
        cell.width = Inches(3.25)
        set_cell_margins(cell, 60, 60, 60, 60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(EVIDENCE / filename), width=Inches(2.95))
        cp = cell.add_paragraph(label)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cp.runs:
            run.italic = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(92, 107, 122)
    doc.add_paragraph()


def page_break(doc):
    doc.add_page_break()


def create_document():
    build_er_diagram()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 45, 58)

    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(29)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title_style.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 18), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(7)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("CampusFix  |  Actividad Evaluable 06  |  Pagina ")
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(92, 107, 122)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    r = p.add_run("CAMPUSFIX")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0, 156, 149)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    p.add_run("Modelado e implementacion de la base de datos cloud")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Actividad Evaluable 06")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(23, 54, 93)

    doc.add_paragraph()
    add_table(doc, ["Campo", "Informacion"], [
        ("Programa", "Ingenieria de Sistemas"),
        ("Asignatura", "Mobile Cloud Computing - IS0249 - 210"),
        ("Estudiante", "Jeyson Arley Sanchez Gomez"),
        ("Codigo", "202210063611"),
        ("Producto", "CampusFix"),
        ("Fecha de entrega", "14 de octubre de 2026"),
    ], widths=[1.75, 4.85], font_size=10.2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run("Documento de evidencia tecnica")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(92, 107, 122)

    page_break(doc)

    # Executive summary and model
    add_heading(doc, "1 Resumen de la implementacion")
    add_body(doc, "CampusFix ya utiliza Firebase Authentication para el acceso y Cloud Firestore para registrar, consultar y auditar reportes de incidentes. La prueba funcional realizada el 14 de octubre de 2026 creo el reporte dQSWv596BoEtSkaPVgPt, genero su evento inicial en la subcoleccion history y permitio consultar el detalle desde la aplicacion.")
    add_body(doc, "El modelo separa identidad, catalogos, operacion e historial en cuatro estructuras: users, categories, reports y reports/{reportId}/history. Las reglas publicadas limitan el acceso por identidad y rol, mientras cuatro indices compuestos respaldan las consultas principales.")

    add_heading(doc, "2 Modelo logico", 1)
    add_table(doc, ["Entidad", "Proposito", "Identificador"], [
        ("users", "Perfil, rol y estado del usuario autenticado.", "UID de Authentication"),
        ("categories", "Catalogo ordenado de tipos de incidente.", "ID semantico"),
        ("reports", "Datos principales y estado actual del incidente.", "ID automatico"),
        ("history", "Eventos inmutables asociados a un reporte.", "ID automatico"),
    ], widths=[1.25, 3.75, 1.6])
    add_heading(doc, "Relaciones principales", 2)
    add_bullets(doc, [
        "Un usuario puede crear muchos reportes y cada reporte conserva un unico creador.",
        "Una categoria puede clasificar muchos reportes y cada reporte utiliza una categoria.",
        "Un reporte puede tener muchos eventos de historial dentro de su propia subcoleccion.",
        "Un usuario autenticado actua como autor de cada evento de historial.",
        "La asignacion a personal responsable es opcional y se representa con assigneeId.",
    ])

    page_break(doc)

    # ER diagram
    add_heading(doc, "3 Diagrama de entidad relacion")
    add_body(doc, "Firestore es una base de datos documental y no aplica llaves foraneas fisicas. Las relaciones se representan mediante identificadores, se validan con reglas de seguridad y se mantienen coherentes desde el servicio de datos de la aplicacion.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(DIAGRAM), width=Inches(7.0))
    caption(doc, "Figura 1. Modelo conceptual de CampusFix implementado en Cloud Firestore.")
    add_heading(doc, "Estructura fisica", 2)
    add_table(doc, ["Ruta", "Contenido"], [
        ("users/{uid}", "Perfil y rol del usuario."),
        ("categories/{categoryId}", "Catalogo de categorias."),
        ("reports/{reportId}", "Documento principal del incidente."),
        ("reports/{reportId}/history/{eventId}", "Trazabilidad del reporte."),
    ], widths=[3.0, 3.6])

    page_break(doc)

    # Dictionary 1
    add_heading(doc, "4 Diccionario de datos")
    add_heading(doc, "Coleccion users", 2)
    add_table(doc, ["Campo", "Tipo", "Req", "Restriccion y descripcion"], [
        ("displayName", "string", "Si", "Nombre visible, entre 2 y 80 caracteres."),
        ("email", "string", "Si", "Correo de acceso asociado a Authentication."),
        ("role", "string", "Si", "student, staff o admin; controla autorizacion."),
        ("active", "boolean", "Si", "Permite desactivar sin eliminar trazabilidad."),
        ("createdAt", "timestamp", "Si", "Hora del servidor; inmutable."),
        ("updatedAt", "timestamp", "Si", "Hora del servidor de la ultima actualizacion."),
    ], widths=[1.15, 1.0, 0.55, 3.9], font_size=8.8)

    add_heading(doc, "Coleccion categories", 2)
    add_table(doc, ["Campo", "Tipo", "Req", "Restriccion y descripcion"], [
        ("name", "string", "Si", "Nombre mostrado en el formulario."),
        ("description", "string", "Si", "Alcance funcional de la categoria."),
        ("active", "boolean", "Si", "Indica si puede usarse en nuevos reportes."),
        ("order", "number", "Si", "Entero mayor o igual a cero para ordenar."),
    ], widths=[1.15, 1.0, 0.55, 3.9], font_size=8.8)
    add_body(doc, "Catalogo inicial: infraestructura, tecnologia, seguridad y aseo.", bold_lead="Catalogo inicial:")

    page_break(doc)

    # Dictionary reports
    add_heading(doc, "5 Diccionario de reports")
    add_table(doc, ["Campo", "Tipo", "Req", "Restriccion y descripcion"], [
        ("title", "string", "Si", "Resumen del incidente, entre 5 y 100 caracteres."),
        ("description", "string", "Si", "Detalle del problema, entre 10 y 1000 caracteres."),
        ("categoryId", "string", "Si", "ID de una categoria activa."),
        ("priority", "string", "Si", "Baja, Media o Alta."),
        ("status", "string", "Si", "Inicia en Abierto; luego En proceso o Resuelto."),
        ("location", "map", "Si", "Incluye label y admite coordenadas opcionales."),
        ("reporterId", "string", "Si", "UID autenticado; inmutable."),
        ("assigneeId", "string/null", "Si", "UID del responsable o null."),
        ("evidencePath", "string/null", "Si", "Ruta futura de fotografia o null."),
        ("createdAt", "timestamp", "Si", "Hora de creacion generada por el servicio."),
        ("updatedAt", "timestamp", "Si", "Fecha de ultima modificacion."),
        ("resolvedAt", "timestamp/null", "Si", "Fecha de cierre o null."),
    ], widths=[1.15, 1.0, 0.55, 3.9], font_size=8.3)
    add_heading(doc, "Mapa location", 2)
    add_table(doc, ["Campo", "Tipo", "Restriccion"], [
        ("label", "string", "Ubicacion legible, entre 3 y 150 caracteres."),
        ("latitude", "number/null", "Valor opcional entre -90 y 90."),
        ("longitude", "number/null", "Valor opcional entre -180 y 180."),
    ], widths=[1.5, 1.3, 3.8], font_size=9)

    page_break(doc)

    # History + business rules
    add_heading(doc, "6 Historial integridad y seguridad")
    add_heading(doc, "Subcoleccion history", 2)
    add_table(doc, ["Campo", "Tipo", "Restriccion y descripcion"], [
        ("type", "string", "created, status_changed, assigned o commented."),
        ("actorId", "string", "UID autenticado que genera el evento."),
        ("previousStatus", "string/null", "Estado anterior cuando aplique."),
        ("newStatus", "string/null", "Nuevo estado cuando aplique."),
        ("comment", "string/null", "Nota opcional de hasta 500 caracteres."),
        ("createdAt", "timestamp", "Momento exacto; evento inmutable."),
    ], widths=[1.45, 1.25, 4.0], font_size=8.8)
    add_heading(doc, "Reglas de negocio aplicadas", 2)
    add_bullets(doc, [
        "Todo reporte nuevo inicia en estado Abierto y sin responsable asignado.",
        "El reporterId debe coincidir con el UID de la sesion autenticada.",
        "El estudiante solo consulta sus reportes y no puede eliminarlos.",
        "El personal con rol staff o admin puede gestionar el flujo operativo.",
        "El historial permite crear eventos, pero impide su edicion o eliminacion.",
        "Los estados y prioridades se limitan a los valores definidos en el modelo.",
    ])
    add_body(doc, "La fotografia es opcional en la interfaz. En esta entrega evidencePath permanece en null cuando no se adjunta una imagen; el campo queda preparado para la integracion posterior con Firebase Storage.")

    page_break(doc)

    # Firestore evidence
    add_heading(doc, "7 Evidencia de la implementacion en Firestore")
    add_body(doc, "El proyecto campusfix-corte2-jeyson-2026 contiene documentos para usuarios, categorias y reportes. Las capturas siguientes muestran la estructura implementada y la correspondencia con el modelo.")
    add_image(doc, "02-firestore-categorias.png", 6.9, "Figura 2. Catalogo de categorias almacenado en Cloud Firestore.")
    add_heading(doc, "Documento de reporte", 2)
    add_image(doc, "03-firestore-reporte.png", 6.9, "Figura 3. Documento de reporte con campos, tipos y valores definidos.")

    page_break(doc)

    add_heading(doc, "8 Evidencia de trazabilidad y autenticacion")
    add_image(doc, "04-firestore-historial.png", 6.9, "Figura 4. Subcoleccion history asociada al reporte.")
    add_image(doc, "07-authentication.png", 6.9, "Figura 5. Usuarios registrados mediante correo y contrasena en Firebase Authentication.")

    page_break(doc)

    # Rules and indexes
    add_heading(doc, "9 Reglas e indices")
    add_body(doc, "Las reglas de seguridad fueron publicadas y los indices compuestos quedaron habilitados. Esto permite consultar por propietario, estado, prioridad y orden de categorias sin exponer informacion de otros usuarios.")
    add_image(doc, "05-reglas-publicadas.png", 6.9, "Figura 6. Reglas de seguridad publicadas en Firestore.")
    add_heading(doc, "Indices compuestos", 2)
    add_image(doc, "06-indices-habilitados.png", 6.9, "Figura 7. Cuatro indices compuestos habilitados.")

    page_break(doc)

    # App evidence 1
    add_heading(doc, "10 Prueba funcional desde CampusFix")
    add_body(doc, "La aplicacion se conecto al proyecto Firebase mediante el SDK web modular. El usuario autenticado creo un incidente con categoria, titulo, descripcion, ubicacion y prioridad, y recibio confirmacion con el ID real generado por Firestore.")
    add_two_images(doc, "08-app-inicio.png", "Figura 8. Inicio posterior a la autenticacion.", "09-app-formulario.png", "Figura 9. Datos principales del nuevo reporte.")
    add_two_images(doc, "10-app-evidencia.png", "Figura 10. Ubicacion y prioridad del incidente.", "11-app-confirmacion.png", "Figura 11. Confirmacion con ID de Firestore.")

    page_break(doc)

    # App evidence 2 and conclusions
    add_heading(doc, "11 Consulta y seguimiento")
    add_image(doc, "12-app-detalle.png", 4.05, "Figura 12. Detalle del reporte dQSWv596BoEtSkaPVgPt y seguimiento inicial.")
    add_body(doc, "La pantalla de detalle recupera el mismo identificador, categoria, estado, titulo, descripcion, ubicacion y prioridad guardados en Firestore. El primer evento de seguimiento confirma la recepcion del reporte y demuestra la lectura de datos desde la nube.")

    add_heading(doc, "12 Coherencia tecnica y resultado", 1)
    add_table(doc, ["Criterio", "Resultado verificado"], [
        ("Modelado y ER", "Cuatro entidades, relaciones y estructura documental definidas."),
        ("Diccionario", "Campos, tipos, obligatoriedad, restricciones y proposito documentados."),
        ("Implementacion cloud", "Authentication y Firestore activos con reglas e indices."),
        ("Coherencia", "La app escribe y consulta los mismos campos descritos en el modelo."),
        ("Sustentacion", "Flujo autenticado probado de extremo a extremo con un reporte real."),
    ], widths=[1.65, 4.95], font_size=9)

    add_heading(doc, "Conclusiones", 2)
    add_bullets(doc, [
        "El modelo satisface el registro, la consulta y la trazabilidad de incidentes de CampusFix.",
        "Las reglas protegen la propiedad de los reportes y reservan la gestion operativa para roles autorizados.",
        "Los indices habilitados soportan las consultas previstas sin descargar colecciones completas.",
        "La prueba funcional confirma la persistencia real en la nube y la lectura posterior desde la aplicacion.",
    ])
    add_body(doc, "Como accion administrativa de entrega, el propietario del proyecto puede agregar el correo endiromero@gmail.com en Usuarios y permisos de Firebase si desea conceder acceso directo al docente.")

    doc.core_properties.title = "CampusFix Modelado e implementacion de la base de datos cloud"
    doc.core_properties.subject = "Actividad Evaluable 06 de Mobile Cloud Computing"
    doc.core_properties.author = "Jeyson Arley Sanchez Gomez"
    doc.core_properties.keywords = "CampusFix, Firebase, Firestore, base de datos cloud"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    create_document()
